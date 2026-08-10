"""Read the local column-inventory .xlsx (openpyxl). No SharePoint / M365 auth.

Returns (rows, headers) where rows is a list of {header: cell_value} dicts — the same
shape snowflake_source produces, so inventory.py can map either identically.
"""

import os

from .config import ConfigError


def read_rows(excel_cfg, excel_path):
    if not excel_path:
        raise ConfigError(
            "No Excel path. Pass --excel-path, set EXCEL_PATH, or source.excel.path in config."
        )
    if not os.path.isfile(excel_path):
        raise ConfigError("Excel file not found: {}".format(excel_path))

    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ConfigError(
            "openpyxl is not installed. Run: pip install -r procs/cpnc/requirements.txt"
        )

    sheet_name = excel_cfg.get("sheet", "Sheet1")
    header_row = int(excel_cfg.get("header_row", 1))

    wb = load_workbook(filename=excel_path, data_only=True, read_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ConfigError(
                "Sheet '{}' not found in {}.\nSheets present: {}".format(
                    sheet_name, os.path.basename(excel_path), ", ".join(wb.sheetnames)
                )
            )
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)

        # Advance to the header row (1-based).
        headers = None
        for idx, values in enumerate(rows_iter, start=1):
            if idx == header_row:
                headers = [("" if v is None else str(v).strip()) for v in values]
                break
        if not headers or not any(headers):
            raise ConfigError(
                "No header row found at row {} of sheet '{}'.".format(header_row, sheet_name)
            )

        rows = []
        for values in rows_iter:
            if values is None or all(v is None for v in values):
                continue  # skip fully blank rows
            row = {}
            for i, header in enumerate(headers):
                if not header:
                    continue
                row[header] = values[i] if i < len(values) else None
            rows.append(row)
    finally:
        wb.close()

    return rows, headers
